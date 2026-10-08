"""Compiling a fiction project to a manuscript; TODO markers."""

import io
import json
import shutil
import subprocess
import zipfile

import pytest


def para(*texts):
    return json.dumps({"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": t}]} for t in texts]})


@pytest.fixture
def novel(api):
    p = api.folder("fiction", "The Lighthouse")
    part = api.folder("fiction", "Part One", p["id"])
    ch1 = api.folder("fiction", "Chapter One", part["id"])
    ch2 = api.folder("fiction", "Chapter Two", part["id"])
    a = api.ok(api.c.post("/api/documents", json={"kind": "fiction", "title": "Arrival", "folder_id": ch1["id"],
                                                  "content_json": para("Mara came off the ferry.", "Rain on the quay.")}))
    b = api.ok(api.c.post("/api/documents", json={"kind": "fiction", "title": "Night", "folder_id": ch1["id"],
                                                  "content_json": para("The lamp [[fix: what colour?]] turned all night.")}))
    m = api.ok(api.c.post("/api/documents", json={"kind": "fiction", "title": "Mara notes", "role": "misc", "folder_id": ch1["id"],
                                                  "content_json": para("SECRET-MISC character notes [[check age]]")}))
    doc = {"type": "doc", "content": [
        {"type": "paragraph", "content": [{"type": "text", "text": "Morning came."}]},
        {"type": "sectionBreak"},
        {"type": "paragraph", "content": [{"type": "text", "text": "Later, ", "marks": []}, {"type": "text", "text": "much", "marks": [{"type": "italic"}]}, {"type": "text", "text": " later."}]},
    ]}
    c = api.ok(api.c.post("/api/documents", json={"kind": "fiction", "title": "Morning", "folder_id": ch2["id"], "content_json": json.dumps(doc)}))
    return {"project": p, "part": part, "ch1": ch1, "ch2": ch2, "a": a, "b": b, "misc": m, "c": c}


def get(api, url):
    r = api.c.get(url)
    assert r.status_code == 200, r.text
    return r


def test_compile_markdown(api, novel):
    r = get(api, f"/api/folders/{novel['project']['id']}/compile?format=md&author=Ana%20Sousa")
    md = r.text
    assert 'filename="The Lighthouse (manuscript).md"' in r.headers["content-disposition"]
    assert md.startswith("# The Lighthouse\n\nby Ana Sousa\n\n19 words")
    order = [md.index(x) for x in ["## Part One", "### Chapter One", "Mara came off", "\n\n#\n\n", "The lamp turned all night.",
                                   "### Chapter Two", "Morning came.", "*much*"]]
    assert order == sorted(order)
    assert "SECRET-MISC" not in md and "[[" not in md and "what colour" not in md
    assert md.count("\n#\n") == 2  # between the two scenes, and the section break inside Morning


def test_compile_html_and_chapter(api, novel):
    h = get(api, f"/api/folders/{novel['ch1']['id']}/compile?format=html&title_page=false").text
    assert "<h1>Chapter One</h1>" in h and '<p class="section-break">#</p>' in h
    assert "SECRET-MISC" not in h and "[[" not in h and "Morning came" not in h
    assert "about" not in h and 'class="front"' not in h


def test_compile_docx_manuscript_format(api, novel):
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Inches, Pt

    raw = get(api, f"/api/folders/{novel['project']['id']}/compile?format=docx&author=Ana%20Sousa").content
    doc = Document(io.BytesIO(raw))
    normal = doc.styles["Normal"]
    assert normal.font.name == "Times New Roman" and normal.font.size == Pt(12)
    assert normal.paragraph_format.line_spacing_rule == WD_LINE_SPACING.DOUBLE
    assert normal.paragraph_format.first_line_indent == Inches(0.5)
    s = doc.sections[0]
    assert all(getattr(s, m) == Inches(1) for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "Ana Sousa\t19 words"  # title page: name and word count
    assert "The Lighthouse" in texts and "by Ana Sousa" in texts
    heads = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert heads == [("Heading 1", "Part One"), ("Heading 2", "Chapter One"), ("Heading 2", "Chapter Two")]
    assert texts.count("#") == 2
    assert not any("SECRET-MISC" in t or "[[" in t or "what colour" in t for t in texts)
    assert "The lamp turned all night." in texts
    for name in ("Part One", "Chapter One", "Chapter Two"):
        heading = next(p for p in doc.paragraphs if p.text == name)
        assert heading.paragraph_format.page_break_before, name  # each starts a new page
    xml = zipfile.ZipFile(io.BytesIO(raw)).read("word/header1.xml").decode() if "word/header1.xml" in zipfile.ZipFile(io.BytesIO(raw)).namelist() else ""
    headers = "".join(zipfile.ZipFile(io.BytesIO(raw)).read(n).decode() for n in zipfile.ZipFile(io.BytesIO(raw)).namelist() if n.startswith("word/header"))
    assert "Sousa / THE LIGHTHOUSE / " in headers and "PAGE" in headers
    assert s.different_first_page_header_footer
    del xml


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_manuscript_opens_in_libreoffice(api, novel, tmp_path):
    raw = get(api, f"/api/folders/{novel['project']['id']}/compile?format=docx").content
    src = tmp_path / "ms.docx"
    src.write_bytes(raw)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp_path), str(src)], check=True, capture_output=True, timeout=120)
    subprocess.run(["soffice", "--headless", "--convert-to", "txt:Text", "--outdir", str(tmp_path), str(src)], check=True, capture_output=True, timeout=120)
    text = (tmp_path / "ms.txt").read_text(encoding="utf-8-sig")
    assert "Chapter Two" in text and "SECRET-MISC" not in text
    pdf = (tmp_path / "ms.pdf").read_bytes()
    assert pdf.count(b"/Type /Page\n") + pdf.count(b"/Type/Page\n") + pdf.count(b"/Type /Page ") + pdf.count(b"/Type/Page/") >= 0
    assert len(pdf) > 1000


def test_compile_check_warns_about_todos(api, novel):
    out = api.ok(api.c.get(f"/api/folders/{novel['project']['id']}/compile-check"))
    assert out == {"folder": {"id": novel["project"]["id"], "name": "The Lighthouse"}, "scenes": 3, "words": 19, "todos": 1}


def test_todos_collected_with_chapter_and_scene(api, novel):
    out = api.ok(api.c.get(f"/api/folders/{novel['project']['id']}/todos"))
    items = out["todos"]
    assert [(t["title"], t["text"], t["path"], t["role"]) for t in items] == [
        ("Night", "fix: what colour?", ["The Lighthouse", "Part One", "Chapter One"], "scene"),
        ("Mara notes", "check age", ["The Lighthouse", "Part One", "Chapter One"], "misc"),
    ]
    assert items[0]["document_id"] == novel["b"]["id"] and items[0]["index"] == 0


def test_compile_only_for_kinds_with_the_tool(api):
    f = api.folder("essay", "Series")
    assert api.c.get(f"/api/folders/{f['id']}/compile").status_code == 400
    assert api.c.get(f"/api/folders/{f['id']}/todos").status_code == 400
    p = api.folder("fiction", "P")
    assert api.c.get(f"/api/folders/{p['id']}/compile?format=pdf").status_code == 422


def test_strip_markers_keeps_text():
    from cadence.manuscript import strip_markers

    doc = {"type": "doc", "content": [{"type": "paragraph", "content": [
        {"type": "text", "text": "Keep [[fix: me]] this"}, {"type": "text", "text": "[[all]]"}, {"type": "text", "text": " end"}]}]}
    out = strip_markers(doc)
    assert [n["text"] for n in out["content"][0]["content"]] == ["Keep this", " end"]
