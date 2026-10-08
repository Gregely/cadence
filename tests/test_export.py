import io
import json
import shutil
import subprocess
import zipfile

import pytest

ESSAY = {
    "type": "doc",
    "content": [
        {"type": "paragraph", "content": [
            {"type": "text", "text": "A "},
            {"type": "text", "text": "bold", "marks": [{"type": "bold"}]},
            {"type": "text", "text": " and "},
            {"type": "text", "text": "italic ", "marks": [{"type": "italic"}]},
            {"type": "text", "text": "claim"},
            {"type": "footnote", "attrs": {"text": "See Thoreau, Walking, p. 4.", "sourceId": None}},
            {"type": "text", "text": " with a "},
            {"type": "text", "text": "link", "marks": [{"type": "link", "attrs": {"href": "https://example.org/a?b=1"}}]},
            {"type": "text", "text": " and <script>alert(1)</script> * _stars_."},
        ]},
        {"type": "heading", "attrs": {"level": 2}, "content": [{"type": "text", "text": "Second part"}]},
        {"type": "blockquote", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Quoted line"}]}]},
        {"type": "sectionBreak"},
        {"type": "paragraph", "content": [{"type": "text", "text": "# not a heading"}]},
    ],
}

POEM = {
    "type": "doc",
    "content": [
        {"type": "paragraph", "content": [
            {"type": "text", "text": "grey at the water"}, {"type": "hardBreak"},
            {"type": "text", "text": "\twaiting"}, {"type": "hardBreak"},
            {"type": "text", "text": "    still", "marks": [{"type": "italic"}]},
        ]},
        {"type": "paragraph", "content": [{"type": "text", "text": "second stanza"}]},
    ],
}


def make(api, kind, title, doc, folder_id=None):
    return api.ok(api.c.post("/api/documents", json={"kind": kind, "title": title, "content_json": json.dumps(doc), "folder_id": folder_id}))


def get(api, url):
    r = api.c.get(url)
    assert r.status_code == 200, r.text
    return r


def test_markdown_export(api):
    d = make(api, "essay", "On Walking", ESSAY)
    r = get(api, f"/api/documents/{d['id']}/export?format=md")
    assert r.headers["content-type"].startswith("text/markdown")
    assert 'filename="On Walking.md"' in r.headers["content-disposition"]
    md = r.text
    assert md.startswith("# On Walking\n")
    assert "A **bold** and *italic* claim[^1]" in md
    assert "[link](<https://example.org/a?b=1>)" in md
    assert "\\<script\\>" in md and "\\* \\_stars\\_" in md
    assert "## Second part" in md
    assert "> Quoted line" in md
    assert "* * *" in md
    assert "\\# not a heading" in md
    assert md.rstrip().endswith("[^1]: See Thoreau, Walking, p. 4.")


def test_html_export_is_escaped(api):
    d = make(api, "essay", "On <Walking>", ESSAY)
    h = get(api, f"/api/documents/{d['id']}/export?format=html").text
    assert "<title>On &lt;Walking&gt;</title>" in h
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in h and "<script>" not in h
    assert '<sup id="fnref1"><a href="#fn1">1</a></sup>' in h
    assert '<li id="fn1">See Thoreau, Walking, p. 4.' in h
    assert '<a href="https://example.org/a?b=1">link</a>' in h
    assert "<strong>bold</strong>" in h and "<em>italic </em>" in h
    assert "<blockquote><p>Quoted line</p></blockquote>" in h


def _docx(api, url):
    from docx import Document

    r = get(api, url)
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    return r.content, Document(io.BytesIO(r.content))


def test_docx_export_structure(api):
    d = make(api, "essay", "On Walking", ESSAY)
    raw, doc = _docx(api, f"/api/documents/{d['id']}/export?format=docx")
    paras = [(p.style.name, p.text) for p in doc.paragraphs]
    assert paras[0] == ("Title", "On Walking")
    assert ("Heading 1", "Second part") in paras  # Title takes the top level
    assert ("Quote", "Quoted line") in paras
    assert ("Normal", "* * *") in paras
    body = paras[1][1]
    assert body.startswith("A bold and italic claim")
    runs = doc.paragraphs[1].runs
    assert any(r.bold and r.text == "bold" for r in runs)
    assert any(r.italic and r.text == "italic " for r in runs)
    z = zipfile.ZipFile(io.BytesIO(raw))
    notes = z.read("word/footnotes.xml").decode()
    assert "See Thoreau, Walking, p. 4." in notes
    document_xml = z.read("word/document.xml").decode()
    assert 'w:footnoteReference w:id="1"' in document_xml
    rels = z.read("word/_rels/document.xml.rels").decode()
    assert "footnotes.xml" in rels and "https://example.org/a?b=1" in rels
    assert "footnotes+xml" in z.read("[Content_Types].xml").decode()


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_docx_opens_in_libreoffice(api, tmp_path):
    d = make(api, "essay", "On Walking", ESSAY)
    raw, _ = _docx(api, f"/api/documents/{d['id']}/export?format=docx")
    src = tmp_path / "walking.docx"
    src.write_bytes(raw)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "html", "--outdir", str(tmp_path), str(src)],
        check=True, capture_output=True, timeout=120,
    )
    text = (tmp_path / "walking.html").read_text(encoding="utf-8", errors="replace")
    assert "Second part" in text and "Quoted line" in text
    assert "See Thoreau, Walking, p. 4." in text  # footnote rendered by Word-compatible reader


def test_poetry_export_keeps_lines_and_indentation(api):
    d = make(api, "poetry", "Heron", POEM)
    md = get(api, f"/api/documents/{d['id']}/export?format=md").text
    assert "grey at the water\\\n    waiting\\\n*    still*" in md or \
        "grey at the water\\\n    waiting\\\n    *still*" in md
    assert "\n\nsecond stanza" in md
    h = get(api, f"/api/documents/{d['id']}/export?format=html").text
    assert '<div class="poem"><p>grey at the water<br>\twaiting<br><em>    still</em></p>' in h
    raw, doc = _docx(api, f"/api/documents/{d['id']}/export?format=docx")
    xml = zipfile.ZipFile(io.BytesIO(raw)).read("word/document.xml").decode()
    assert "<w:tab/>" in xml and "<w:br/>" in xml
    assert 'xml:space="preserve">    still' in xml


def test_folder_exports_as_one_document_in_order(api):
    novel = api.folder("fiction", "The Novel")
    part1 = api.folder("fiction", "Part One", novel["id"])
    part2 = api.folder("fiction", "Part Two", novel["id"])
    s = lambda t: {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": t}]}]}
    make(api, "fiction", "Arrival", s("She came by train."), part1["id"])
    make(api, "fiction", "", s("Untitled scene text."), part1["id"])
    make(api, "fiction", "Departure", s("He left at dawn."), part2["id"])
    make(api, "fiction", "Epilogue", s("Years later."), novel["id"])
    # Reorder: move Part Two before Part One.
    api.ok(api.c.post(f"/api/folders/{part2['id']}/move", json={"parent_id": novel["id"], "index": 0}))
    md = get(api, f"/api/folders/{novel['id']}/export?format=md").text
    order = [md.index(x) for x in ["# The Novel", "## Part Two", "### Departure", "## Part One", "### Arrival",
                                   "She came by train.", "* * *", "Untitled scene text.", "## Epilogue", "Years later."]]
    assert order == sorted(order), md
    raw, doc = _docx(api, f"/api/folders/{novel['id']}/export?format=docx")
    headings = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith(("Heading", "Title"))]
    assert headings == [("Title", "The Novel"), ("Heading 1", "Part Two"), ("Heading 2", "Departure"),
                        ("Heading 1", "Part One"), ("Heading 2", "Arrival"), ("Heading 1", "Epilogue")]
    # Fiction paragraphs after the first get a first-line indent.
    h = get(api, f"/api/folders/{novel['id']}/export?format=html").text
    assert '<div class="indented">' in h


def test_trashed_documents_left_out_of_folder_export(api):
    f = api.folder("essay", "Series")
    keep = make(api, "essay", "Keep", ESSAY, f["id"])
    gone = make(api, "essay", "Gone", ESSAY, f["id"])
    api.ok(api.c.delete(f"/api/documents/{gone['id']}"))
    md = get(api, f"/api/folders/{f['id']}/export?format=md").text
    assert "## Keep" in md and "Gone" not in md
    assert keep


def test_export_rejected_for_diary_and_bad_format(api, diary):
    d = diary("secret")
    for fmt in ("md", "html", "docx"):
        r = api.c.get(f"/api/documents/{d['id']}/export?format={fmt}")
        assert r.status_code == 400
    e = make(api, "essay", "x", ESSAY)
    assert api.c.get(f"/api/documents/{e['id']}/export?format=pdf").status_code == 422


def test_full_backup_zip(api):
    f = api.folder("essay", "Series: one/two")
    make(api, "essay", "Walking", ESSAY, f["id"])
    make(api, "poetry", "Heron", POEM)
    trashed = make(api, "essay", "Trashed", ESSAY)
    api.ok(api.c.delete(f"/api/documents/{trashed['id']}"))
    r = get(api, "/api/export/full")
    assert r.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(r.content))
    names = set(z.namelist())
    assert {"cadence.sqlite3", "data.json", "README.txt"} <= names
    assert "documents/essay/Series one two/001 Walking.md" in names
    assert "documents/poetry/001 Heron.md" in names
    assert not any("Trashed" in n for n in names)
    data = json.loads(z.read("data.json"))
    assert {d["title"] for d in data["documents"]} == {"Walking", "Heron", "Trashed"}


def test_safe_filename():
    from cadence.export import safe_filename

    assert safe_filename('a/b\\c:d*e?"f<g>h|') == "a b c d e f g h"
    assert safe_filename("CON") == "cadence-CON"
    assert safe_filename("...") == "untitled"
    assert len(safe_filename("x" * 300)) == 80
