"""Export documents to Markdown, HTML and Word (.docx).

All three renderers walk the stored ProseMirror JSON. A folder exports as one
combined document in display order (subfolders first, then documents, as in
the library), with headings shifted under the document titles.

Word files are written with python-docx. Footnotes become real Word
footnotes: python-docx has no API for them, so ``_DocxFootnotes`` adds the
footnotes part itself.
"""

from __future__ import annotations

import html
import io
import json
import re
import sqlite3
from dataclasses import dataclass, field
from typing import Any, Callable

from .errors import Forbidden, Invalid
from .kinds import Kind, get_kind
from .library import doc_row, folder_row, require_kind

FORMATS = {
    "md": ("text/markdown; charset=utf-8", "md"),
    "html": ("text/html; charset=utf-8", "html"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
}

SAFE_HREF = re.compile(r"^(https?:|mailto:|#)", re.I)


# ---------------------------------------------------------------- model


@dataclass
class Part:
    """One document in an export, with the heading level of its title."""

    title: str
    doc: dict
    kind: Kind
    title_level: int  # 0 = no title heading
    first: bool = True


@dataclass
class Section:
    """A folder heading inside a combined export."""

    title: str
    level: int


@dataclass
class Source:
    id: int
    title: str
    author: str = ""
    url: str = ""
    published: str = ""

    def citation(self) -> str:
        bits = [b for b in (self.author, f"“{self.title}”" if self.title else "", self.published, self.url) if b]
        return ", ".join(bits)


@dataclass
class Export:
    title: str
    kind: Kind
    items: list[Part | Section] = field(default_factory=list)
    sources: dict[int, Source] = field(default_factory=dict)
    single: bool = False  # one document: its title is the export title
    scene_break: str = "* * *"  # between untitled parts and for section breaks
    front: list[str] = field(default_factory=list)  # lines under the title (compile)


def _sources_lookup(conn: sqlite3.Connection) -> dict[int, Source]:
    """Sources (Stage 5 research table) if present."""
    try:
        rows = conn.execute("SELECT id, title, author, url, published FROM sources").fetchall()
    except sqlite3.OperationalError:
        return {}
    return {r["id"]: Source(r["id"], r["title"] or "", r["author"] or "", r["url"] or "", r["published"] or "") for r in rows}


def _doc_json(row: sqlite3.Row) -> dict:
    try:
        return json.loads(row["content_json"])
    except ValueError:
        return {"type": "doc", "content": []}


def _exportable(kind: Kind) -> None:
    if not kind.exportable or kind.encrypted:
        raise Forbidden(f"{kind.label} cannot be exported from the server")


def document_export(conn: sqlite3.Connection, doc_id: int) -> Export:
    row = doc_row(conn, doc_id)
    kind = require_kind(row["kind"])
    _exportable(kind)
    title = row["title"] or "Untitled"
    exp = Export(title=title, kind=kind, sources=_sources_lookup(conn), single=True)
    exp.items.append(Part(title=title, doc=_doc_json(row), kind=kind, title_level=1))
    return exp


def folder_export(conn: sqlite3.Connection, folder_id: int) -> Export:
    """A folder as one document, in library order."""
    top = folder_row(conn, folder_id)
    kind = require_kind(top["kind"])
    _exportable(kind)
    exp = Export(title=top["name"], kind=kind, sources=_sources_lookup(conn))

    def walk(fid: int, depth: int) -> None:
        # depth 0 is the exported folder itself (its name is the title, h1).
        for sub in conn.execute(
            "SELECT id, name FROM folders WHERE parent_id = ? AND deleted_at IS NULL ORDER BY sort_order, id", (fid,)
        ).fetchall():
            exp.items.append(Section(sub["name"], min(6, 2 + depth)))
            walk(sub["id"], depth + 1)
        first = True
        for d in conn.execute(
            "SELECT * FROM documents WHERE folder_id = ? AND deleted_at IS NULL ORDER BY sort_order, id", (fid,)
        ).fetchall():
            title = d["title"].strip()
            exp.items.append(
                Part(title=title, doc=_doc_json(d), kind=kind, title_level=min(6, 2 + depth) if title else 0, first=first)
            )
            first = False

    walk(folder_id, 0)
    return exp


# ---------------------------------------------------------------- shared helpers


def _children(node: dict) -> list[dict]:
    content = node.get("content")
    return content if isinstance(content, list) else []


def _attrs(node: dict) -> dict:
    a = node.get("attrs")
    return a if isinstance(a, dict) else {}


def _marks(node: dict) -> list[dict]:
    m = node.get("marks")
    return [x for x in m if isinstance(x, dict)] if isinstance(m, list) else []


def _heading_level(part: Part, level: Any) -> int:
    level = level if isinstance(level, int) else 2
    if part.title_level:
        return min(6, part.title_level + (level - 1))
    return min(6, level)


class Footnotes:
    """Collects footnotes in order, numbering across the whole export."""

    def __init__(self, sources: dict[int, Source]):
        self.notes: list[str] = []
        self.sources = sources
        self.cited: list[int] = []

    def add(self, attrs: dict) -> int:
        text = str(attrs.get("text") or "").strip()
        sid = attrs.get("sourceId")
        if isinstance(sid, int) and sid in self.sources:
            self.cite(sid)
            if not text:
                text = self.sources[sid].citation()
        self.notes.append(text)
        return len(self.notes)

    def cite(self, sid: Any) -> None:
        if isinstance(sid, int) and sid in self.sources and sid not in self.cited:
            self.cited.append(sid)


# ---------------------------------------------------------------- markdown


MD_ESCAPE = re.compile(r"([\\`*_\[\]<>|])")


def _md_text(text: str) -> str:
    return MD_ESCAPE.sub(r"\\\1", text)


def _md_line_start(line: str) -> str:
    """Escape things that would turn a paragraph line into another block."""
    if re.match(r"^(#{1,6}\s|>|[-+*]\s|\d+[.)]\s|={3,}|-{3,}|~{3,})", line):
        line = "\\" + line if not re.match(r"^\d", line) else re.sub(r"^(\d+)([.)])", r"\1\\\2", line)
    return line


def _md_preserve_indent(line: str) -> str:
    """Leading spaces/tabs survive Markdown as non-breaking spaces."""
    m = re.match(r"^[ \t]+", line)
    if not m:
        return line
    lead = m.group(0).replace("\t", "    ")
    return " " * len(lead) + line[m.end():]


class MarkdownRenderer:
    def __init__(self, exp: Export):
        self.exp = exp
        self.notes = Footnotes(exp.sources)

    def inline(self, nodes: list[dict], poetry: bool) -> str:
        out: list[str] = []
        for n in nodes:
            t = n.get("type")
            if t == "text":
                text = _md_text(n.get("text", ""))
                marks = {m.get("type"): m for m in _marks(n)}
                if "citation" in marks:
                    self.notes.cite(_attrs(marks["citation"]).get("sourceId"))
                core = text.strip()
                if core:
                    lead = text[: len(text) - len(text.lstrip())]
                    trail = text[len(text.rstrip()):]
                    if "italic" in marks:
                        core = f"*{core}*"
                    if "bold" in marks:
                        core = f"**{core}**"
                    if "link" in marks:
                        href = str(_attrs(marks["link"]).get("href") or "")
                        if SAFE_HREF.match(href):
                            core = f"[{core}](<{href}>)"
                    text = lead + core + trail
                out.append(text)
            elif t == "hardBreak":
                out.append("\\\n")
            elif t == "footnote":
                out.append(f"[^{self.notes.add(_attrs(n))}]")
        text = "".join(out)
        lines = text.split("\n")
        lines = [_md_line_start(_md_preserve_indent(l) if poetry else l) for l in lines]
        return "\n".join(lines)

    def blocks(self, nodes: list[dict], part: Part, prefix: str = "") -> list[str]:
        out: list[str] = []
        poetry = "poetryLines" in part.kind.extensions
        for n in nodes:
            t = n.get("type")
            if t == "paragraph":
                out.append(self.inline(_children(n), poetry))
            elif t == "heading":
                level = _heading_level(part, _attrs(n).get("level"))
                out.append("#" * level + " " + self.inline(_children(n), False).replace("\\\n", " "))
            elif t == "blockquote":
                inner = "\n\n".join(self.blocks(_children(n), part))
                out.append("\n".join(("> " + l) if l else ">" for l in inner.split("\n")))
            elif t in ("bulletList", "orderedList"):
                items = []
                for i, li in enumerate(_children(n), 1):
                    marker = "- " if t == "bulletList" else f"{i}. "
                    body = "\n\n".join(self.blocks(_children(li), part))
                    lines = body.split("\n")
                    items.append(marker + lines[0] + "".join("\n" + " " * len(marker) + l if l else "\n" for l in lines[1:]))
                out.append("\n".join(items))
            elif t == "sectionBreak":
                out.append(self.exp.scene_break)
        return out

    def render(self) -> str:
        chunks = [f"# {_md_text(self.exp.title)}"]
        chunks.extend(_md_text(line) for line in self.exp.front)
        single = self.exp.single
        for item in self.exp.items:
            if isinstance(item, Section):
                chunks.append("#" * item.level + " " + _md_text(item.title))
                continue
            if item.title_level and not single:
                chunks.append("#" * item.title_level + " " + _md_text(item.title))
            elif not item.title_level and not item.first:
                chunks.append(self.exp.scene_break)
            chunks.extend(b for b in self.blocks(_children(item.doc), item) if b.strip())
        if self.notes.cited:
            chunks.append("## Sources")
            chunks.append("\n".join(f"- {_md_text(self.exp.sources[s].citation())}" for s in self.notes.cited))
        if self.notes.notes:
            chunks.append("\n".join(f"[^{i}]: {_md_text(t)}" for i, t in enumerate(self.notes.notes, 1)))
        return "\n\n".join(chunks) + "\n"


# ---------------------------------------------------------------- html


HTML_STYLE = """
body { font-family: Georgia, 'Iowan Old Style', serif; max-width: 38rem; margin: 3rem auto; padding: 0 1.25rem;
       line-height: 1.6; color: #1f1d1a; background: #fff; }
h1, h2, h3, h4 { line-height: 1.3; }
blockquote { margin: 1em 0; padding-left: 1em; border-left: 2px solid #999; }
.section-break { text-align: center; letter-spacing: .5em; margin: 1.5em 0; }
.poem p { white-space: pre-wrap; tab-size: 4; }
.indented p + p { text-indent: 1.5em; margin-top: 0; }
.indented p { margin-bottom: 0; }
sup a { text-decoration: none; }
.footnotes { border-top: 1px solid #999; margin-top: 3rem; font-size: .9em; }
"""


class HtmlRenderer:
    def __init__(self, exp: Export):
        self.exp = exp
        self.notes = Footnotes(exp.sources)

    def inline(self, nodes: list[dict]) -> str:
        out: list[str] = []
        for n in nodes:
            t = n.get("type")
            if t == "text":
                text = html.escape(n.get("text", ""), quote=False)
                for m in _marks(n):
                    mt = m.get("type")
                    if mt == "italic":
                        text = f"<em>{text}</em>"
                    elif mt == "bold":
                        text = f"<strong>{text}</strong>"
                    elif mt == "link":
                        href = str(_attrs(m).get("href") or "")
                        if SAFE_HREF.match(href):
                            text = f'<a href="{html.escape(href)}">{text}</a>'
                    elif mt == "citation":
                        self.notes.cite(_attrs(m).get("sourceId"))
                        text = f"<cite>{text}</cite>"
                out.append(text)
            elif t == "hardBreak":
                out.append("<br>")
            elif t == "footnote":
                i = self.notes.add(_attrs(n))
                out.append(f'<sup id="fnref{i}"><a href="#fn{i}">{i}</a></sup>')
        return "".join(out)

    def blocks(self, nodes: list[dict], part: Part) -> list[str]:
        out = []
        for n in nodes:
            t = n.get("type")
            if t == "paragraph":
                out.append(f"<p>{self.inline(_children(n))}</p>")
            elif t == "heading":
                lv = _heading_level(part, _attrs(n).get("level"))
                out.append(f"<h{lv}>{self.inline(_children(n))}</h{lv}>")
            elif t == "blockquote":
                out.append("<blockquote>" + "".join(self.blocks(_children(n), part)) + "</blockquote>")
            elif t in ("bulletList", "orderedList"):
                tag = "ul" if t == "bulletList" else "ol"
                items = "".join("<li>" + "".join(self.blocks(_children(li), part)) + "</li>" for li in _children(n))
                out.append(f"<{tag}>{items}</{tag}>")
            elif t == "sectionBreak":
                out.append(f'<p class="section-break">{html.escape(self.exp.scene_break)}</p>')
        return out

    def render(self) -> str:
        body = [f"<h1>{html.escape(self.exp.title)}</h1>"]
        body.extend(f'<p class="front">{html.escape(line)}</p>' for line in self.exp.front)
        single = self.exp.single
        css_class = "poem" if "poetryLines" in self.exp.kind.extensions else (
            "indented" if self.exp.kind.theme.get("text_indent", "0") not in ("0", "") else "")
        for item in self.exp.items:
            if isinstance(item, Section):
                body.append(f"<h{item.level}>{html.escape(item.title)}</h{item.level}>")
                continue
            if item.title_level and not single:
                body.append(f"<h{item.title_level}>{html.escape(item.title)}</h{item.title_level}>")
            elif not item.title_level and not item.first:
                body.append(f'<p class="section-break">{html.escape(self.exp.scene_break)}</p>')
            body.append(f'<div class="{css_class}">' + "\n".join(self.blocks(_children(item.doc), item)) + "</div>")
        if self.notes.cited:
            body.append("<h2>Sources</h2><ul>" + "".join(
                f"<li>{html.escape(self.exp.sources[s].citation())}</li>" for s in self.notes.cited) + "</ul>")
        if self.notes.notes:
            body.append('<section class="footnotes"><ol>' + "".join(
                f'<li id="fn{i}">{html.escape(t)} <a href="#fnref{i}">↩</a></li>'
                for i, t in enumerate(self.notes.notes, 1)) + "</ol></section>")
        return (
            "<!doctype html>\n<html lang=\"en\"><head><meta charset=\"utf-8\">"
            "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">"
            f"<title>{html.escape(self.exp.title)}</title><style>{HTML_STYLE}</style></head>\n<body>\n"
            + "\n".join(body)
            + "\n</body></html>\n"
        )


# ---------------------------------------------------------------- docx


def _docx_modules():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part as OpcPart
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Cm, Pt

    return Document, WD_ALIGN_PARAGRAPH, RT, PackURI, OpcPart, parse_xml, nsdecls, qn, Cm, Pt


class _DocxFootnotes:
    """Adds a footnotes part to a python-docx document and real footnote references."""

    CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"

    def __init__(self, document):
        _, _, RT, PackURI, OpcPart, parse_xml, nsdecls, qn, _, _ = _docx_modules()
        self.qn = qn
        self.parse_xml = parse_xml
        self.nsdecls = nsdecls
        self.root = parse_xml(
            f'<w:footnotes {nsdecls("w")}>'
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
            '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
            "</w:footnotes>"
        )
        self.document = document
        self.count = 0
        self.part = None
        self._OpcPart = OpcPart
        self._PackURI = PackURI
        self._RT = RT

    def add(self, paragraph, text: str) -> None:
        qn = self.qn
        self.count += 1
        fid = str(self.count)
        run = paragraph.add_run()
        rpr = run._r.get_or_add_rPr()
        rpr.append(self.parse_xml(f'<w:vertAlign {self.nsdecls("w")} w:val="superscript"/>'))
        ref = self.parse_xml(f'<w:footnoteReference {self.nsdecls("w")} w:id="{fid}"/>')
        run._r.append(ref)
        note = self.parse_xml(
            f'<w:footnote {self.nsdecls("w")} w:id="{fid}"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr>'
            '<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r>'
            '<w:r><w:t xml:space="preserve"> </w:t></w:r></w:p></w:footnote>'
        )
        t = note.makeelement(qn("w:t"), {})
        t.text = text
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r = note.makeelement(qn("w:r"), {})
        r.append(t)
        note.find(qn("w:p")).append(r)
        self.root.append(note)

    def finish(self) -> None:
        if not self.count:
            return
        from lxml import etree

        blob = etree.tostring(self.root, xml_declaration=True, encoding="UTF-8", standalone=True)
        part = self._OpcPart(self._PackURI("/word/footnotes.xml"), self.CT, blob, self.document.part.package)
        self.document.part.relate_to(part, self._RT.FOOTNOTES)


class DocxRenderer:
    def __init__(self, exp: Export):
        self.exp = exp
        self.notes = Footnotes(exp.sources)

    def render(self) -> bytes:
        Document, ALIGN, RT, _, _, parse_xml, nsdecls, qn, Cm, Pt = _docx_modules()
        self.ALIGN, self.RT, self.parse_xml, self.nsdecls, self.qn, self.Cm = ALIGN, RT, parse_xml, nsdecls, qn, Cm
        doc = Document()
        self.doc = doc
        self.fn = _DocxFootnotes(doc)
        normal = doc.styles["Normal"]
        normal.font.name = "Georgia"
        normal.font.size = Pt(12)
        doc.core_properties.title = self.exp.title
        doc.core_properties.author = ""
        doc.core_properties.comments = ""
        doc.add_heading(self.exp.title, level=0)
        single = self.exp.single
        for item in self.exp.items:
            if isinstance(item, Section):
                doc.add_heading(item.title, level=min(9, item.level - 1))
                continue
            if item.title_level and not single:
                doc.add_heading(item.title, level=min(9, item.title_level - 1))
            elif not item.title_level and not item.first:
                self.section_break()
            self.indent_next = False
            self.blocks(_children(item.doc), item, None)
        if self.notes.cited:
            doc.add_heading("Sources", level=1)
            for s in self.notes.cited:
                doc.add_paragraph(self.exp.sources[s].citation(), style="List Bullet")
        self.fn.finish()
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def section_break(self) -> None:
        p = self.doc.add_paragraph("* * *")
        p.alignment = self.ALIGN.CENTER
        self.indent_next = False

    def blocks(self, nodes: list[dict], part: Part, style: str | None, list_style: str | None = None) -> None:
        indent = part.kind.theme.get("text_indent", "0") not in ("0", "")
        for n in nodes:
            t = n.get("type")
            if t == "paragraph":
                p = self.doc.add_paragraph(style=list_style or style)
                if indent and self.indent_next and not (list_style or style):
                    p.paragraph_format.first_line_indent = self.Cm(0.75)
                if indent:
                    p.paragraph_format.space_after = 0
                self.inline(p, _children(n))
                self.indent_next = True
            elif t == "heading":
                level = _heading_level(part, _attrs(n).get("level"))
                p = self.doc.add_heading(level=min(9, level - 1))
                self.inline(p, _children(n))
                self.indent_next = False
            elif t == "blockquote":
                self.blocks(_children(n), part, "Quote")
                self.indent_next = False
            elif t in ("bulletList", "orderedList"):
                ls = "List Bullet" if t == "bulletList" else "List Number"
                for li in _children(n):
                    self.blocks(_children(li), part, style, ls)
                self.indent_next = False
            elif t == "sectionBreak":
                self.section_break()

    def inline(self, p, nodes: list[dict]) -> None:
        for n in nodes:
            t = n.get("type")
            if t == "text":
                marks = {m.get("type"): m for m in _marks(n)}
                text = n.get("text", "")
                if "citation" in marks:
                    self.notes.cite(_attrs(marks["citation"]).get("sourceId"))
                href = str(_attrs(marks["link"]).get("href") or "") if "link" in marks else ""
                if href and SAFE_HREF.match(href) and not href.startswith("#"):
                    self.hyperlink(p, text, href, "bold" in marks, "italic" in marks)
                    continue
                run = p.add_run(text)
                run.bold = "bold" in marks or None
                run.italic = "italic" in marks or None
            elif t == "hardBreak":
                p.add_run().add_break()
            elif t == "footnote":
                attrs = _attrs(n)
                self.notes.add(attrs)
                self.fn.add(p, self.notes.notes[-1])

    def hyperlink(self, p, text: str, href: str, bold: bool, italic: bool) -> None:
        qn = self.qn
        r_id = p.part.relate_to(href, self.RT.HYPERLINK, is_external=True)
        link = p._p.makeelement(qn("w:hyperlink"), {qn("r:id"): r_id})
        run = p._p.makeelement(qn("w:r"), {})
        rpr = run.makeelement(qn("w:rPr"), {})
        style = rpr.makeelement(qn("w:rStyle"), {qn("w:val"): "Hyperlink"})
        rpr.append(style)
        underline = rpr.makeelement(qn("w:u"), {qn("w:val"): "single"})
        rpr.append(underline)
        if bold:
            rpr.append(rpr.makeelement(qn("w:b"), {}))
        if italic:
            rpr.append(rpr.makeelement(qn("w:i"), {}))
        run.append(rpr)
        t = run.makeelement(qn("w:t"), {})
        t.text = text
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        run.append(t)
        link.append(run)
        p._p.append(link)


# ---------------------------------------------------------------- entry points


RENDERERS: dict[str, Callable[[Export], Any]] = {
    "md": lambda e: MarkdownRenderer(e).render().encode("utf-8"),
    "html": lambda e: HtmlRenderer(e).render().encode("utf-8"),
    "docx": lambda e: DocxRenderer(e).render(),
}


def render(exp: Export, fmt: str) -> tuple[bytes, str, str]:
    """Return (bytes, media type, filename)."""
    if fmt not in FORMATS:
        raise Invalid("format must be md, html or docx")
    media, ext = FORMATS[fmt]
    return RENDERERS[fmt](exp), media, f"{safe_filename(exp.title)}.{ext}"


WINDOWS_RESERVED = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}


def safe_filename(name: str, limit: int = 80) -> str:
    name = re.sub(r'[\x00-\x1f<>:"/\\|?*]+', " ", name)
    name = " ".join(name.split()).strip(" .")
    name = name[:limit].strip(" .")
    if not name or name.upper().split(".")[0] in WINDOWS_RESERVED:
        name = f"cadence-{name}" if name else "untitled"
    return name


def content_disposition(filename: str) -> str:
    from urllib.parse import quote

    ascii_name = filename.encode("ascii", "replace").decode().replace("?", "_").replace('"', "")
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"


def markdown_for_backup(conn: sqlite3.Connection, row: sqlite3.Row, sources: dict[int, Source]) -> str | None:
    kind = get_kind(row["kind"])
    if kind is None or kind.encrypted or not kind.exportable:
        return None
    title = row["title"] or "Untitled"
    exp = Export(title=title, kind=kind, sources=sources, single=True)
    exp.items.append(Part(title=title, doc=_doc_json(row), kind=kind, title_level=1))
    return MarkdownRenderer(exp).render()

